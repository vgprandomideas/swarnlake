import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def main() -> int:
    if len(sys.argv) < 4:
        raise SystemExit("Usage: export_xlsx.py <input_json> <output_xlsx> <title>")

    input_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])
    title = sys.argv[3]

    rows = json.loads(input_path.read_text(encoding="utf-8"))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31] or "Registrations"

    headers = list(rows[0].keys()) if rows else ["Message"]
    sheet.append(headers)

    if rows:
        for row in rows:
            sheet.append([row.get(header, "") for header in headers])
    else:
        sheet.append(["No registrations found"])

    header_fill = PatternFill(fill_type="solid", fgColor="123B59")
    header_font = Font(color="FFFFFF", bold=True)

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for column_index, column_cells in enumerate(sheet.columns, start=1):
        max_length = 0
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, len(value))
        sheet.column_dimensions[get_column_letter(column_index)].width = min(max(max_length + 2, 14), 42)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
